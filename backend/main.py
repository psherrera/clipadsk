"""
Clipadsk - Backend (FastAPI)

Descarga y transcripción de video/audio para periodistas.
- Subtítulos de YouTube → Groq Whisper → Whisper local (en cascada)
- Herramientas periodísticas con IA (resumen, citas con tiempos, datos, ángulos, diarización)

Módulos:
  config.py      variables de entorno, rutas, logging
  text_utils.py  funciones puras (subtítulos, SRT, citas, URLs) — con tests
  media.py       FFmpeg y transcripción (Groq / Whisper local)
  ai.py          limpieza, traducción y análisis con Groq
"""
import os
import re
import io
import json
import time
import uuid
import hmac
import base64
import random
import shutil
import sqlite3
import asyncio
import zipfile
import tempfile
import functools
import subprocess
import sys
import http.cookiejar
from pathlib import Path
from typing import Optional, Any, List
from urllib.parse import quote
from concurrent.futures import ThreadPoolExecutor

import requests
import yt_dlp
from fastapi import FastAPI, Request, HTTPException, BackgroundTasks, UploadFile, File, Form, Response
from fastapi.responses import JSONResponse, FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from config import (
    logger, BASE_DIR, ROOT_DIR, FRONTEND_DIR, DOWNLOAD_FOLDER, DB_FILE, CACHE_FILE,
    HOST, PORT, ADMIN_TOKEN, ALLOWED_ORIGINS, HAS_FFMPEG, FFMPEG_BIN, IS_RENDER,
    MAX_WORKERS, MAX_CONCURRENT_JOBS, MAX_UPLOAD_MB, GROQ_MAX_UPLOAD_MB,
)
from text_utils import (
    sanitize_url, parse_subtitles_to_segments, generate_srt_from_segments, remove_repetitions,
    find_segment_times_for_quote, host_matches, is_allowed_thumbnail_url, safe_filename,
    YOUTUBE_DOMAINS, INSTAGRAM_DOMAINS, TIKTOK_DOMAINS, TWITTER_DOMAINS, FACEBOOK_DOMAINS,
)
import media
import ai

try:
    import instaloader
except ImportError:
    instaloader = None

try:
    from cachetools import TTLCache
except ImportError:  # pragma: no cover
    TTLCache = None


app = FastAPI(title="Clipadsk API")

# ─── EJECUCIÓN EN SEGUNDO PLANO ──────────────────────────────────────────────
# Todo lo que bloquea (yt-dlp, FFmpeg, Groq, requests) se corre en hilos para que el
# servidor siga respondiendo (progreso, logs) mientras trabaja.
EXECUTOR = ThreadPoolExecutor(max_workers=MAX_WORKERS)
HEAVY_JOBS = asyncio.Semaphore(MAX_CONCURRENT_JOBS)


async def run_blocking(fn: Any, *args, **kwargs):
    """Corre una función bloqueante en el pool, limitando trabajos pesados simultáneos."""
    async with HEAVY_JOBS:
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(EXECUTOR, functools.partial(fn, *args, **kwargs))


async def run_light(fn: Any, *args, **kwargs):
    """Corre una función bloqueante liviana (HTTP corto, disco) sin ocupar un cupo pesado."""
    return await asyncio.to_thread(fn, *args, **kwargs)


# ─── SEGURIDAD ───────────────────────────────────────────────────────────────
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


@app.middleware("http")
async def block_foreign_origins(request: Request, call_next):
    """
    CORS no frena los POST "simples" (formularios) de otras páginas: igual llegan al
    servidor. Rechazamos cualquier petición a /api/ que venga con un Origin ajeno.
    """
    if request.url.path.startswith("/api/") and request.method not in ("GET", "HEAD", "OPTIONS"):
        origin = request.headers.get("origin")
        if origin and "*" not in ALLOWED_ORIGINS and origin not in ALLOWED_ORIGINS:
            # Mismo host que sirve la app (p. ej. acceso por IP en la red con HOST=0.0.0.0)
            if origin.split("://", 1)[-1] != request.headers.get("host", ""):
                logger.warning(f"Petición bloqueada desde origen no permitido: {origin}")
                return JSONResponse(status_code=403, content={"detail": "Origen no permitido."})
    return await call_next(request)


def require_admin(request: Request):
    """Si ADMIN_TOKEN está definido, exige el header X-ADMIN-TOKEN."""
    if ADMIN_TOKEN:
        provided = request.headers.get('X-ADMIN-TOKEN', '')
        if not hmac.compare_digest(provided, ADMIN_TOKEN):
            raise HTTPException(status_code=403, detail="Se requiere token de administrador para esta operación.")


# ─── PROGRESO, LOGS Y RESULTADOS (memoria, con vencimiento) ──────────────────
if TTLCache:
    progress_store = TTLCache(maxsize=500, ttl=7200)
    log_store = TTLCache(maxsize=500, ttl=7200)
    result_store = TTLCache(maxsize=100, ttl=3600)
else:  # pragma: no cover
    progress_store, log_store, result_store = {}, {}, {}


def update_progress(uid: Optional[str], progress: int, text: str):
    if uid:
        progress_store[uid] = {"progress": progress, "text": text}
        add_log(uid, f"Progreso {progress}%: {text}")


def add_log(uid: Optional[str], message: str):
    logger.debug(f"[{uid}] {message}")
    if not uid:
        return
    log_store.setdefault(uid, []).append(f"[{time.strftime('%H:%M:%S')}] {message}")


def store_result(uid: Optional[str], result: dict):
    if uid:
        result_store[uid] = result


@app.get("/api/progress/{uid}")
async def get_progress(uid: str):
    return progress_store.get(uid, {"progress": 0, "text": "Procesando en el servidor..."})


@app.get("/api/logs/{uid}")
async def get_logs(uid: str):
    return {"logs": "\n".join(log_store.get(uid, ["No hay logs disponibles para esta sesion."]))}


@app.get("/api/result/{uid}")
async def get_transcript_result(uid: str):
    """Recupera una transcripción terminada si la conexión del navegador se cortó."""
    result = result_store.get(uid)
    if result is None:
        raise HTTPException(status_code=404, detail="Resultado no disponible. La transcripción puede no haber terminado o el UID es incorrecto.")
    return result


# ─── CACHÉ DE TRANSCRIPCIONES (SQLite) ───────────────────────────────────────

def db_connect():
    return sqlite3.connect(DB_FILE, timeout=10)


def init_db():
    with db_connect() as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS transcripts
                        (url TEXT PRIMARY KEY, transcript TEXT, date_added TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
        if os.path.exists(CACHE_FILE):  # migración del formato JSON viejo
            logger.info("Migrando historial de JSON a SQLite...")
            try:
                with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                    for url, text in json.load(f).items():
                        conn.execute("INSERT OR IGNORE INTO transcripts (url, transcript) VALUES (?, ?)", (url, text))
                conn.commit()
                os.replace(CACHE_FILE, CACHE_FILE + ".migrated")
            except Exception:
                logger.exception("Error en migración de cache JSON a SQLite")


def cache_get(key: str) -> Optional[dict]:
    try:
        with db_connect() as conn:
            row = conn.execute("SELECT transcript FROM transcripts WHERE url = ?", (key,)).fetchone()
    except Exception:
        logger.exception("Error leyendo cache")
        return None
    if not row:
        return None
    try:
        parsed = json.loads(row[0])
        if isinstance(parsed, dict) and "transcript" in parsed:
            return parsed
    except (TypeError, ValueError):
        pass
    return {"transcript": row[0], "srt": "", "segments": []}  # entradas viejas (solo texto)


def cache_set(key: str, data: dict):
    try:
        with db_connect() as conn:
            conn.execute("INSERT OR REPLACE INTO transcripts (url, transcript) VALUES (?, ?)", (key, json.dumps(data)))
    except Exception:
        logger.exception("Error guardando en cache")


init_db()


# ─── YT-DLP ──────────────────────────────────────────────────────────────────

USER_AGENTS = [
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36',
    'Mozilla/5.0 (Macintosh; Intel Mac OS X 14_3_1) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36',
    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36 Edg/124.0.0.0',
]
DESKTOP_UA = USER_AGENTS[0]
MOBILE_UA = 'Mozilla/5.0 (iPhone; CPU iPhone OS 17_3_1 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.3.1 Mobile/15E148 Safari/604.1'

_env_cookie_files: dict = {}


def _cookie_file_from_env(var_names) -> Optional[str]:
    """Cookies en base64 por variable de entorno (Render/Docker). Se escriben una sola vez."""
    for var in var_names:
        b64 = os.environ.get(var)
        if not b64:
            continue
        if var not in _env_cookie_files:
            try:
                tmp = tempfile.NamedTemporaryFile(mode='w', suffix='.txt', delete=False, encoding='utf-8')
                tmp.write(base64.b64decode(b64).decode())
                tmp.close()
                _env_cookie_files[var] = tmp.name
            except Exception:
                logger.exception(f"Cookies inválidas en {var}")
                continue
        return _env_cookie_files[var]
    return None


def find_cookie_file(url: str) -> Optional[str]:
    is_instagram = host_matches(url, INSTAGRAM_DOMAINS)
    env_vars = ['INSTAGRAM_COOKIES_B64', 'COOKIES_B64'] if is_instagram else ['COOKIES_B64']
    env_file = _cookie_file_from_env(env_vars)
    if env_file:
        return env_file
    names = ['cookies_ig.txt', 'cookies.txt'] if is_instagram else ['cookies.txt']
    candidates = [os.environ.get('COOKIES_PATH', '')]
    for name in names:
        candidates += [f'/etc/secrets/{name}', os.path.join(BASE_DIR, name), os.path.join(ROOT_DIR, name)]
    return next((p for p in candidates if p and os.path.exists(p)), None)


def get_robust_opts(target_url: str, extra: Optional[dict] = None) -> dict:
    """Opciones base de yt-dlp: cookies, user-agent y headers por plataforma."""
    opts = {
        'quiet': True,
        'no_warnings': False,
        'cachedir': False,
        'noplaylist': True,
        'nocheckcertificate': True,
        'ignoreerrors': False,
        'user_agent': random.choice(USER_AGENTS),
        'js_runtimes': {'deno': {}, 'node': {}},
        'remote_components': ['ejs:github'],
    }
    if FFMPEG_BIN:
        opts['ffmpeg_location'] = FFMPEG_BIN
    cookie = find_cookie_file(target_url)
    if cookie:
        opts['cookiefile'] = cookie

    if host_matches(target_url, TIKTOK_DOMAINS):
        opts['user_agent'] = MOBILE_UA
        opts['http_headers'] = {'Referer': 'https://www.tiktok.com/', 'Accept-Language': 'es-419,es;q=0.9,en;q=0.8'}
    elif host_matches(target_url, YOUTUBE_DOMAINS + TWITTER_DOMAINS + FACEBOOK_DOMAINS) or 'mediadelivery.net' in target_url:
        opts['user_agent'] = DESKTOP_UA
    elif 'b-cdn.net' in target_url and 'playlist.m3u8' in target_url:
        referer = os.environ.get('BUNNY_REFERER', '')
        if referer:
            opts['http_headers'] = {'Referer': referer}

    opts.update(extra or {})
    return opts


def _strategies(url: str) -> list:
    """Estrategias en orden: normal (con cookies), móvil sin cookies, solo iOS."""
    def no_cookies(clients):
        def apply(opts):
            opts.pop('cookiefile', None)
            opts['extractor_args'] = {'youtube': {'player_client': clients}}
        return apply
    strategies = [("normal", lambda o: None)]
    if host_matches(url, YOUTUBE_DOMAINS):
        strategies += [("móvil sin cookies", no_cookies(['android', 'ios'])), ("solo iOS", no_cookies(['ios']))]
    return strategies


def ytdlp_extract(url: str, extra: Optional[dict] = None) -> dict:
    """extract_info con reintentos (bloqueante)."""
    errors = []
    for name, tweak in _strategies(url):
        opts = get_robust_opts(url, extra)
        tweak(opts)
        try:
            with yt_dlp.YoutubeDL(opts) as ydl:
                info = ydl.extract_info(url, download=False)
            if info:
                return info
        except Exception as e:
            errors.append(f"{name}: {str(e)[:150]}")
            logger.debug(f"extract_info ({name}) falló: {e}")
    raise RuntimeError(" | ".join(errors) or "Sin información")


def ytdlp_download(url: str, extra: dict, done_check, on_attempt=None):
    """
    Descarga con reintentos (bloqueante). done_check() decide si el intento
    produjo lo esperado (yt-dlp a veces "termina bien" sin bajar nada).
    """
    errors = []
    for n, (name, tweak) in enumerate(_strategies(url), start=1):
        if on_attempt:
            on_attempt(n, name)
        opts = get_robust_opts(url, extra)
        tweak(opts)
        try:
            with yt_dlp.YoutubeDL(opts) as ydl:
                ydl.download([url])
        except Exception as e:
            errors.append(f"{name}: {str(e)[:150]}")
            logger.debug(f"Descarga ({name}) falló: {e}")
        if done_check():
            return
    raise RuntimeError(" | ".join(errors) or "yt-dlp no descargó ningún archivo")


# ─── BUNNY STREAM (videos embebidos en páginas web) ──────────────────────────

KNOWN_PLATFORMS = YOUTUBE_DOMAINS + INSTAGRAM_DOMAINS + TIKTOK_DOMAINS + TWITTER_DOMAINS + FACEBOOK_DOMAINS


def detect_bunny_embed(page_url: str) -> Optional[str]:
    """Busca un iframe de Bunny Stream o un m3u8 de Bunny CDN en una página web."""
    if 'mediadelivery.net' in page_url or 'b-cdn.net' in page_url:
        return None
    if host_matches(page_url, KNOWN_PLATFORMS) or not page_url.startswith(('http://', 'https://')):
        return None
    try:
        cookies = {}
        cookie_path = os.path.join(BASE_DIR, 'cookies.txt')
        if os.path.exists(cookie_path):
            try:
                cj = http.cookiejar.MozillaCookieJar(cookie_path)
                cj.load(ignore_discard=True, ignore_expires=True)
                cookies = {c.name: c.value for c in cj}
            except Exception:
                pass
        resp = requests.get(page_url, headers={'User-Agent': DESKTOP_UA, 'Accept-Language': 'es-419,es;q=0.9,en;q=0.8'},
                            cookies=cookies, timeout=15)
        if resp.status_code != 200:
            return None
        html = resp.text
        m = re.search(r'(https?://iframe\.mediadelivery\.net/embed/\d+/[a-f0-9\-]+)', html, re.I)
        if m:
            return m.group(1)
        m = re.search(r'mediadelivery\.net[^"\']*?/embed/(\d+)/([a-f0-9\-]{36})', html, re.I)
        if m:
            return f"https://iframe.mediadelivery.net/embed/{m.group(1)}/{m.group(2)}"
        m = re.search(r'(https?://[a-z0-9\-]+\.b-cdn\.net/[a-f0-9\-]+/playlist\.m3u8)', html, re.I)
        if m:
            return m.group(1)
    except Exception as e:
        logger.debug(f"detect_bunny_embed error para {page_url}: {e}")
    return None


async def resolve_url(raw_url: str) -> str:
    url = sanitize_url(raw_url)
    if not url.startswith(('http://', 'https://')):
        raise HTTPException(status_code=400, detail="La URL debe empezar con http:// o https://")
    bunny = await run_light(detect_bunny_embed, url)
    if bunny:
        logger.info(f"Bunny embed encontrado en {url} → {bunny}")
        return bunny
    return url


# ─── INSTAGRAM ───────────────────────────────────────────────────────────────

def get_instagram_info(url: str) -> dict:
    """Info de un Reel/Post de Instagram con instaloader (bloqueante)."""
    if not instaloader:
        raise RuntimeError("instaloader no está instalado")
    L = instaloader.Instaloader(download_videos=True, download_video_thumbnails=True, download_geotags=False,
                                download_comments=False, save_metadata=False, compress_json=False, quiet=True)
    ig_user, ig_pass = os.environ.get('IG_USER', ''), os.environ.get('IG_PASS', '')
    if ig_user and ig_pass:
        try:
            L.login(ig_user, ig_pass)
        except Exception as e:
            logger.debug(f"Instaloader login falló: {e}")
    match = re.search(r'/(reel|reels|p|tv)/([A-Za-z0-9_-]+)', url)
    if not match:
        raise RuntimeError("No se pudo extraer el código del post de Instagram")
    shortcode = match.group(2)
    post = instaloader.Post.from_shortcode(L.context, shortcode)
    try:
        thumbnail = post.url
    except Exception:
        thumbnail = None
    return {
        'shortcode': shortcode,
        'title': post.caption[:100] if post.caption else f"Instagram Reel {shortcode}",
        'thumbnail': thumbnail,
        'duration': int(post.video_duration) if post.is_video and post.video_duration else None,
        'uploader': post.owner_username,
        'is_video': post.is_video,
        'video_url': post.video_url if post.is_video else None,
    }


def get_instagram_carousel_info(url: str, cookies_path: Optional[str] = None) -> Optional[list]:
    """Lista de imágenes/videos de un carrusel usando gallery-dl (bloqueante)."""
    candidates = [
        os.path.join(os.path.dirname(sys.executable), "gallery-dl.exe"),
        os.path.join(os.path.dirname(sys.executable), "gallery-dl"),
        os.path.join(BASE_DIR, "venv", "Scripts", "gallery-dl.exe"),
    ]
    gallery_dl = next((c for c in candidates if os.path.exists(c)), None) or shutil.which("gallery-dl")
    if not gallery_dl:
        logger.warning("gallery-dl no está instalado (pip install gallery-dl)")
        return None
    cmd = [gallery_dl, "-j"]
    if cookies_path and os.path.exists(cookies_path):
        cmd += ["--cookies", cookies_path]
    cmd.append(url)
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', timeout=120)
        if res.returncode != 0:
            logger.warning(f"gallery-dl falló ({res.returncode}): {res.stderr[:300]}")
            return None
        return [{"url": item[1], "metadata": item[2]} for item in json.loads(res.stdout)
                if isinstance(item, list) and len(item) >= 3 and item[0] == 3]
    except Exception:
        logger.exception("Error extrayendo carrusel con gallery-dl")
        return None


def proxied_thumbnail(url: Optional[str]) -> Optional[str]:
    return f"/api/proxy-thumbnail?url={quote(url, safe='')}" if url else None


def http_download(url: str, dest: str, headers: dict, timeout: int = 60):
    with requests.get(url, headers=headers, stream=True, timeout=timeout) as r:
        r.raise_for_status()
        with open(dest, 'wb') as f:
            for chunk in r.iter_content(chunk_size=65536):
                f.write(chunk)


# ─── MODELOS ─────────────────────────────────────────────────────────────────

class VideoRequest(BaseModel):
    url: str
    format_id: Optional[str] = "best"
    start_time: Optional[str] = None
    end_time: Optional[str] = None
    groq_api_key: Optional[str] = None
    uid: Optional[str] = None
    target_lang: Optional[str] = "es"  # es, en, original


class ChatRequest(BaseModel):
    url: Optional[str] = ""
    question: str
    transcript: str
    groq_api_key: Optional[str] = None


class AnalyzeRequest(BaseModel):
    transcript: str
    mode: str  # summary | data | angle | diarization
    groq_api_key: Optional[str] = None


class QuotesRequest(BaseModel):
    transcript: str
    segments: list = []
    groq_api_key: Optional[str] = None


class ExportDocxRequest(BaseModel):
    title: str
    uploader: Optional[str] = ""
    url: Optional[str] = ""
    description: Optional[str] = ""
    transcript: str


def require_groq(api_key: Optional[str]):
    client = ai.get_groq_client(api_key)
    if not client:
        raise HTTPException(status_code=503, detail="Groq API no configurada. Cargá tu API Key en Configuración.")
    return client


def ai_error(e: Exception) -> HTTPException:
    if isinstance(e, ai.RateLimitedError) or ai.is_rate_limit(e):
        return HTTPException(status_code=429, detail="⚠️ Límite de uso de Groq alcanzado. Podés:\n1. Esperar unos minutos e intentar de nuevo.\n2. Configurar tu propia API key de Groq en Configuración (gratis en console.groq.com).")
    logger.exception("Error de IA")
    return HTTPException(status_code=500, detail=f"Error al analizar: {e}")


# ─── ENDPOINTS: SALUD ────────────────────────────────────────────────────────

@app.get("/api/health")
async def health():
    """Chequeo rápido (no sale a internet)."""
    return {"status": "ok", "ffmpeg": HAS_FFMPEG, "whisper_local": media.WHISPER_MODEL_AVAILABLE,
            "groq_server_key": ai.get_groq_client() is not None, "yt_dlp": yt_dlp.version.__version__}


@app.get("/api/health/cookies")
async def check_cookies():
    """Verifica si las cookies de YouTube siguen funcionando (prueba real, tarda unos segundos)."""
    test_url = "https://www.youtube.com/watch?v=dQw4w9WgXcQ"
    try:
        info = await run_light(lambda: yt_dlp.YoutubeDL(get_robust_opts(test_url)).extract_info(test_url, download=False))
        return {"status": "ok", "cookie_valid": True, "video_title": info.get('title'),
                "server_time": time.strftime("%Y-%m-%d %H:%M:%S")}
    except Exception as e:
        return {"status": "error", "cookie_valid": False, "error": str(e)[:300],
                "server_time": time.strftime("%Y-%m-%d %H:%M:%S")}


# ─── ENDPOINTS: INFO DE VIDEO ────────────────────────────────────────────────

def _resolution_label(height: Optional[int], fmt: dict) -> str:
    if height:
        for min_h, label in ((2160, "2160p (4K UHD)"), (1440, "1440p (2K QHD)"), (1080, "1080p (Full HD)"),
                             (720, "720p (HD)"), (480, "480p (SD)"), (360, "360p (SD)")):
            if height >= min_h:
                return label
        return f"{height}p"
    note = fmt.get('format_note')
    if note and not re.search(r'\d+x\d+', note) and len(note) < 15:
        return note
    return "Calidad estándar"


MP3_FORMAT = {'format_id': 'mp3', 'ext': 'mp3', 'resolution': 'Solo audio', 'filesize': None, 'label': 'Solo audio (.mp3)'}
BEST_FORMAT = {'format_id': 'best', 'ext': 'mp4', 'resolution': 'Mejor calidad', 'filesize': None, 'label': 'Mejor calidad (.mp4)'}


@app.post("/api/video-info")
async def get_video_info(req: VideoRequest):
    url = await resolve_url(req.url)
    is_instagram = host_matches(url, INSTAGRAM_DOMAINS)

    if is_instagram and instaloader:
        try:
            ig = await run_light(get_instagram_info, url)
            formats = [BEST_FORMAT, MP3_FORMAT] if ig['is_video'] else [
                {'format_id': 'best', 'ext': 'jpg', 'resolution': 'Imagen original', 'filesize': None, 'label': 'Imagen original (.jpg)'}]
            thumb = proxied_thumbnail(ig.get('thumbnail'))
            return {'title': ig['title'], 'thumbnail': thumb, 'max_res_thumbnail': thumb, 'duration': ig.get('duration'),
                    'uploader': ig.get('uploader') or 'Instagram', 'description': ig['title'], 'formats': formats,
                    'has_ffmpeg': HAS_FFMPEG, 'has_subtitles': False}
        except Exception as e:
            logger.debug(f"Instaloader falló, intentando con yt-dlp: {e}")

    try:
        info = await run_blocking(ytdlp_extract, url)
    except Exception as e:
        last_error = str(e)
        logger.error(f"No se pudo obtener info de {url}: {last_error[:300]}")
        if is_instagram:
            files = await run_blocking(get_instagram_carousel_info, url, find_cookie_file(url))
            if files:
                meta = files[0].get("metadata", {})
                title = meta.get("description") or f"Instagram Post {meta.get('post_shortcode', '')}"
                thumb = proxied_thumbnail(files[0].get("url"))
                return {'title': title[:100] + ("..." if len(title) > 100 else ""), 'thumbnail': thumb,
                        'max_res_thumbnail': thumb, 'duration': None,
                        'uploader': meta.get("username") or "Instagram User", 'description': meta.get("description") or "",
                        'formats': [{'format_id': 'carousel_images', 'ext': 'zip', 'resolution': 'Imágenes (ZIP)',
                                     'filesize': None, 'label': f"Conjunto de {len(files)} imágenes (.zip)"}],
                        'has_ffmpeg': HAS_FFMPEG, 'has_subtitles': False, 'can_transcribe': True}
            if "No video formats found" in last_error:
                raise HTTPException(status_code=400, detail="Este post de Instagram no contiene video (es una publicación de fotos). Clipadsk solo puede descargar o transcribir videos y audios.")
        raise HTTPException(status_code=400, detail=f"No pudimos procesar este video. Puede ser privado o la plataforma bloqueó la conexión. Errores: {last_error[:200]}")

    formats, seen = [], set()
    useful = sorted((f for f in info.get('formats') or [] if f.get('vcodec') != 'none'),
                    key=lambda f: f.get('height') or 0, reverse=True)
    for f in useful:
        height = f.get('height')
        res_str = f.get('resolution') or ''
        if not height and 'x' in res_str:
            try:
                height = int(res_str.split('x')[1])
            except (ValueError, IndexError):
                pass
        res = _resolution_label(height, f)
        ext = f.get('ext', 'mp4')
        if f"{res}_{ext}" in seen:
            continue
        seen.add(f"{res}_{ext}")
        formats.append({'format_id': f.get('format_id'), 'ext': ext, 'resolution': res,
                        'filesize': f.get('filesize') or f.get('filesize_approx'), 'label': f"{res} (.{ext})"})
    if not formats:
        formats.append(BEST_FORMAT)
    if HAS_FFMPEG:
        formats.append(MP3_FORMAT)

    thumbnail = info.get('thumbnail')
    if is_instagram:
        thumbnail = proxied_thumbnail(thumbnail)
    return {
        'title': info.get('title'),
        'thumbnail': thumbnail,
        'max_res_thumbnail': thumbnail,
        'duration': info.get('duration'),
        'uploader': info.get('uploader') or "Desconocido",
        'description': (info.get('description') or 'Sin descripción')[:200] + '...',
        'formats': formats,
        'has_ffmpeg': HAS_FFMPEG,
        'has_subtitles': bool(info.get('subtitles') or info.get('automatic_captions')),
    }


# ─── ENDPOINTS: TRANSCRIPCIÓN DE URL ─────────────────────────────────────────

def _find_file(folder: str, prefix: str, contains=()) -> Optional[str]:
    for f in sorted(os.listdir(folder)):
        if f.startswith(prefix) and (not contains or any(c in f for c in contains)):
            return os.path.join(folder, f)
    return None


def transcribe_from_subtitles(url: str, lang: str, tmpdir: str, uid: Optional[str]) -> Optional[dict]:
    """Intenta usar los subtítulos de YouTube (sin descargar audio). Bloqueante."""
    def has_subs():
        return any(f.startswith('sub.') for f in os.listdir(tmpdir))

    extra = {'skip_download': True, 'writesubtitles': True, 'writeautomaticsub': True,
             'subtitleslangs': ['es.*', 'en.*'], 'subtitlesformat': 'vtt/srt/best',
             'outtmpl': os.path.join(tmpdir, 'sub.%(ext)s')}
    try:
        ytdlp_download(url, extra, has_subs,
                       on_attempt=lambda n, name: update_progress(uid, 5 + n * 5, f"Buscando subtítulos ({name})..."))
    except Exception as e:
        add_log(uid, f"Sin subtítulos: {str(e)[:200]}")
        return None

    sub_file = _find_file(tmpdir, 'sub.', ('.es',))
    is_english = False
    if not sub_file:
        sub_file = _find_file(tmpdir, 'sub.', ('.en',))
        is_english = sub_file is not None
    if not sub_file:
        return None

    with open(sub_file, 'r', encoding='utf-8', errors='replace') as f:
        segments = parse_subtitles_to_segments(f.read())
    if not segments:
        return None

    # Los subtítulos automáticos repiten texto entre segmentos: limpiar
    for seg in segments:
        seg["text"] = remove_repetitions(seg["text"])
    if is_english and lang == "es":
        update_progress(uid, 50, "Traduciendo subtítulos al español...")
        for seg, t in zip(segments, ai.translate_texts([s["text"] for s in segments], "es")):
            seg["text"] = t
    raw_text = remove_repetitions(' '.join(s["text"] for s in segments))
    return {"raw": raw_text, "segments": segments, "method": "subtitles"}


def transcribe_from_audio(url: str, lang: str, tmpdir: str, uid: Optional[str], client) -> dict:
    """Descarga el audio y lo transcribe (Groq → Whisper local). Bloqueante."""
    add_log(uid, "Descargando audio para transcribir...")
    extra = {'format': 'bestaudio/best', 'outtmpl': os.path.join(tmpdir, 'audio.%(ext)s')}

    def has_audio():
        return _find_file(tmpdir, 'audio.') is not None

    ytdlp_download(url, extra, has_audio,
                   on_attempt=lambda n, name: update_progress(uid, 20 + n * 3, f"Descargando audio ({name})..."))
    audio_file = _find_file(tmpdir, 'audio.')
    if HAS_FFMPEG:
        update_progress(uid, 30, "Preparando audio...")
        audio_file = media.to_speech_mp3(audio_file, os.path.join(tmpdir, 'speech.mp3'))

    text, segments, method = media.transcribe_audio(
        audio_file, lang, client, tmpdir,
        progress=lambda p, t: update_progress(uid, p, t), log=lambda m: add_log(uid, m))
    return {"raw": remove_repetitions(text.strip()), "segments": segments, "method": method}


def transcribe_instagram_carousel(files: list, client, uid: Optional[str]) -> str:
    """OCR de un carrusel de imágenes con Groq Vision. Bloqueante."""
    parts = []
    for idx, item in enumerate(files):
        update_progress(uid, int(idx / len(files) * 90) + 5, f"Analizando diapositiva {idx + 1}/{len(files)}...")
        try:
            r = requests.get(item["url"], headers={'User-Agent': DESKTOP_UA}, timeout=20)
            r.raise_for_status()
            text = ai.ocr_image(client, base64.b64encode(r.content).decode('utf-8'))
        except Exception as e:
            logger.warning(f"OCR diapositiva {idx + 1} falló: {e}")
            text = f"[Error de extracción: {e}]"
        parts.append(f"--- DIAPOSITIVA {idx + 1} ---\n{text}")
    return "\n\n".join(parts)


def _only_images(files: list) -> bool:
    for f in files:
        meta = f.get("metadata", {})
        if meta.get("video_url") or (meta.get("extension") or "").lower() in ('mp4', 'mov', 'avi', 'mkv', 'webm'):
            return False
    return True


@app.post("/api/transcript")
async def get_transcript(req: VideoRequest):
    uid = req.uid
    lang = req.target_lang or "es"
    try:
        url = await resolve_url(req.url)
    except HTTPException as e:
        return JSONResponse(status_code=e.status_code, content={"error": e.detail})
    add_log(uid, f"Iniciando transcripción para: {url} | Idioma: {lang}")

    cache_key = f"{url}_{lang}"
    cached = await run_light(cache_get, cache_key)
    if cached:
        add_log(uid, "Resultado recuperado de la caché local.")
        return {"transcript": cached.get("transcript", ""), "srt": cached.get("srt", ""),
                "segments": cached.get("segments", []), "method": "cache"}

    client = ai.get_groq_client(req.groq_api_key)
    warnings: List[str] = []
    try:
        # Carrusel de Instagram solo con imágenes → OCR
        if host_matches(url, INSTAGRAM_DOMAINS):
            files = await run_blocking(get_instagram_carousel_info, url, find_cookie_file(url))
            if files and _only_images(files):
                if not client:
                    return JSONResponse(status_code=400, content={"error": "Para extraer texto de imágenes (OCR) necesitás configurar la API Key de Groq en Configuración."})
                add_log(uid, f"Carrusel de {len(files)} imágenes: OCR con Groq Vision")
                text = await run_blocking(transcribe_instagram_carousel, files, client, uid)
                payload = {"transcript": text, "srt": "", "segments": [], "method": "groq_vision_ocr"}
                await run_light(cache_set, cache_key, payload)
                update_progress(uid, 100, "¡Extracción de texto completa!")
                store_result(uid, payload)
                return payload

        with tempfile.TemporaryDirectory() as tmpdir:
            result = None
            if host_matches(url, YOUTUBE_DOMAINS):
                add_log(uid, "Buscando subtítulos de YouTube...")
                result = await run_blocking(transcribe_from_subtitles, url, lang, tmpdir, uid)
            if not result:
                result = await run_blocking(transcribe_from_audio, url, lang, tmpdir, uid, client)

        text = result["raw"]
        if client:
            update_progress(uid, 85, f"Aplicando puntuación y párrafos con IA ({lang})...")
            text = await run_blocking(ai.cleanup_transcript, text, client, lang, warnings)

        segments = result["segments"]
        payload = {"transcript": text, "srt": generate_srt_from_segments(segments), "segments": segments,
                   "method": result["method"]}
        if warnings:
            payload["warnings"] = warnings  # no se cachea: la próxima vez se reintenta la limpieza
        else:
            await run_light(cache_set, cache_key, payload)
        update_progress(uid, 100, "¡Transcripción lista!")
        add_log(uid, f"Transcripción completada ({result['method']}).")
        store_result(uid, payload)
        return payload
    except Exception as e:
        logger.exception("Error en transcripción")
        add_log(uid, f"Error: {e}")
        return JSONResponse(status_code=500, content={"error": str(e)[:500]})


# ─── ENDPOINTS: TRANSCRIPCIÓN DE ARCHIVO ─────────────────────────────────────

ALLOWED_AUDIO_EXTENSIONS = {'.ogg', '.opus', '.mp3', '.m4a', '.wav', '.mp4', '.aac', '.weba', '.webm',
                            '.mov', '.avi', '.mkv', '.flac'}
VIDEO_EXTS = {'.mp4', '.mov', '.avi', '.mkv', '.webm'}


async def save_upload(file: UploadFile, dest: str, max_mb: int) -> float:
    """Guarda el archivo subido en disco por partes (sin cargarlo entero en memoria)."""
    size = 0
    limit = max_mb * 1024 * 1024
    with open(dest, 'wb') as out:
        while True:
            chunk = await file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > limit:
                raise HTTPException(status_code=413, detail=f"El archivo es demasiado grande. Máximo: {max_mb} MB.")
            await run_light(out.write, chunk)
    return size / (1024 * 1024)


@app.post("/api/transcript-file")
async def transcript_audio_file(
    file: UploadFile = File(...),
    target_lang: str = Form(default="es"),
    uid: str = Form(default=None),
    groq_api_key: str = Form(default=None),
    is_local_video: Optional[str] = Form(default=None),
):
    """Transcribe un archivo subido (WhatsApp .ogg/.opus, grabaciones, videos)."""
    client = ai.get_groq_client(groq_api_key)
    if not client and not media.WHISPER_MODEL_AVAILABLE:
        raise HTTPException(status_code=503, detail="Groq API no configurada. Configurá tu API Key en la interfaz o instalá faster-whisper para transcribir localmente.")

    ext = os.path.splitext(file.filename or '')[1].lower()
    if ext not in ALLOWED_AUDIO_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"Formato no soportado: '{ext}'. Formatos válidos: {', '.join(sorted(ALLOWED_AUDIO_EXTENSIONS))}")

    warnings: List[str] = []
    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = os.path.join(tmpdir, f"input{ext}")
        size_mb = await save_upload(file, input_path, MAX_UPLOAD_MB)
        update_progress(uid, 5, "Archivo recibido en el servidor...")
        add_log(uid, f"Archivo recibido: {file.filename} ({size_mb:.2f} MB)")

        try:
            audio_path = input_path
            if HAS_FFMPEG or media.AudioSegment:
                update_progress(uid, 10, "Convirtiendo audio...")
                try:
                    audio_path = await run_blocking(media.to_speech_mp3, input_path, os.path.join(tmpdir, "speech.mp3"))
                except Exception as e:
                    logger.warning(f"Conversión falló, se usa el archivo original: {e}")
            converted_mb = os.path.getsize(audio_path) / (1024 * 1024)
            if converted_mb >= GROQ_MAX_UPLOAD_MB and not HAS_FFMPEG and not media.WHISPER_MODEL_AVAILABLE:
                raise HTTPException(status_code=400, detail="El archivo es demasiado grande para Groq y no se puede dividir porque FFmpeg no está instalado. Instalá FFmpeg o subí un audio más corto.")

            update_progress(uid, 20, "Iniciando transcripción del audio...")
            text, segments, method = await run_blocking(
                media.transcribe_audio, audio_path, target_lang, client, tmpdir,
                progress=lambda p, t: update_progress(uid, p, t), log=lambda m: add_log(uid, m))
            text = remove_repetitions(text.strip())

            if client:
                update_progress(uid, 80, "Aplicando puntuación y párrafos con IA...")
                text = await run_blocking(ai.cleanup_transcript, text, client, target_lang, warnings)

            update_progress(uid, 100, "Completado")
            add_log(uid, "Transcripción de archivo completada.")
            payload = {"transcript": text, "srt": generate_srt_from_segments(segments), "segments": segments,
                       "method": method, "filename": file.filename, "size_mb": round(size_mb, 2)}
            if warnings:
                payload["warnings"] = warnings
            store_result(uid, payload)
            return payload
        except HTTPException:
            raise
        except Exception as e:
            logger.exception("Error transcribiendo archivo")
            add_log(uid, f"Error en transcripción de archivo: {e}")
            raise HTTPException(status_code=500, detail=f"Error al transcribir: {e}")


# ─── ENDPOINTS: DESCARGA ─────────────────────────────────────────────────────

def _remove_later(background_tasks: BackgroundTasks, *paths):
    def remove():
        for p in paths:
            try:
                if p and os.path.exists(p):
                    os.remove(p)
            except OSError:
                logger.warning(f"No se pudo borrar {p}")
    background_tasks.add_task(remove)


def _build_carousel_zip(files: list, zip_path: str):
    with zipfile.ZipFile(zip_path, 'w') as zf:
        for idx, item in enumerate(files):
            ext = item.get("metadata", {}).get("extension") or "jpg"
            try:
                r = requests.get(item["url"], headers={'User-Agent': DESKTOP_UA}, timeout=30)
                if r.status_code == 200:
                    zf.writestr(f"imagen_{idx + 1}.{ext}", r.content)
            except Exception as e:
                logger.warning(f"Error descargando imagen {idx + 1} para zip: {e}")


def _instagram_download(ig: dict, uid: str, as_mp3: bool) -> str:
    tmp_mp4 = os.path.join(DOWNLOAD_FOLDER, f'instagram_{uid}.mp4')
    http_download(ig['video_url'], tmp_mp4, {'User-Agent': MOBILE_UA})
    if not as_mp3:
        return tmp_mp4
    mp3_path = os.path.join(DOWNLOAD_FOLDER, f'instagram_{uid}.mp3')
    try:
        media.run_ffmpeg(['-i', tmp_mp4, '-vn', '-acodec', 'libmp3lame', '-q:a', '2', mp3_path])
    finally:
        os.remove(tmp_mp4)
    return mp3_path


@app.post("/api/download")
async def download_video(req: VideoRequest, background_tasks: BackgroundTasks):
    url = await resolve_url(req.url)
    format_id = req.format_id
    file_uid = uuid.uuid4().hex

    # Carrusel de imágenes (gallery-dl) → ZIP
    if format_id == 'carousel_images':
        files = await run_blocking(get_instagram_carousel_info, url, find_cookie_file(url))
        if not files:
            raise HTTPException(status_code=400, detail="No se pudieron extraer las imágenes del carrusel.")
        zip_path = os.path.join(DOWNLOAD_FOLDER, f"instagram_carousel_{file_uid}.zip")
        await run_blocking(_build_carousel_zip, files, zip_path)
        _remove_later(background_tasks, zip_path)
        user = safe_filename(files[0]["metadata"].get("username", "instagram"), 30)
        return FileResponse(zip_path, filename=f"carrusel_{user}_{file_uid[:6]}.zip", media_type='application/zip')

    # Instagram con instaloader
    if host_matches(url, INSTAGRAM_DOMAINS) and instaloader:
        try:
            ig = await run_light(get_instagram_info, url)
            if not ig['is_video']:
                raise HTTPException(status_code=400, detail="Este post de Instagram no tiene video.")
            as_mp3 = format_id == 'mp3'
            if as_mp3 and not HAS_FFMPEG:
                raise HTTPException(status_code=400, detail="Para descargar como MP3 hace falta FFmpeg.")
            path = await run_blocking(_instagram_download, ig, file_uid, as_mp3)
            _remove_later(background_tasks, path)
            ext = 'mp3' if as_mp3 else 'mp4'
            return FileResponse(path, filename=f"{safe_filename(ig['title'], 30)}_{file_uid[:6]}.{ext}",
                                media_type='audio/mpeg' if as_mp3 else 'video/mp4')
        except HTTPException:
            raise
        except Exception as e:
            logger.debug(f"Instaloader falló, intentando con yt-dlp: {e}")

    output_template = os.path.join(DOWNLOAD_FOLDER, f'%(title).80B_{file_uid}.%(ext)s')
    if format_id == 'mp3':
        if not HAS_FFMPEG:
            raise HTTPException(status_code=400, detail="Para descargar como MP3 hace falta FFmpeg.")
        extra = {'format': 'bestaudio/best', 'outtmpl': output_template,
                 'postprocessors': [{'key': 'FFmpegExtractAudio', 'preferredcodec': 'mp3', 'preferredquality': '192'}]}
    else:
        if not HAS_FFMPEG:
            fmt = 'best[ext=mp4]/best'  # sin FFmpeg no se puede unir video+audio
        elif format_id and format_id not in ('best', 'bestvideo+bestaudio'):
            fmt = f"{format_id}+bestaudio[ext=m4a]/{format_id}+bestaudio/bestvideo[ext=mp4]+bestaudio[ext=m4a]/bestvideo+bestaudio/best[ext=mp4]/best"
        else:
            fmt = 'bestvideo[ext=mp4]+bestaudio[ext=m4a]/bestvideo+bestaudio/best[ext=mp4]/best'
        extra = {'format': fmt, 'outtmpl': output_template, 'merge_output_format': 'mp4'}

    def hook(d):
        if d.get('status') == 'downloading':
            total = d.get('total_bytes') or d.get('total_bytes_estimate')
            if total:
                pct = d.get('downloaded_bytes', 0) / total * 100
                update_progress(req.uid, 20 + int(pct * 0.7), f"Descargando: {pct:.0f}%")
        elif d.get('status') == 'finished':
            update_progress(req.uid, 90, "Descarga completada, procesando con FFmpeg...")
    extra['progress_hooks'] = [hook]

    if req.start_time or req.end_time:
        from yt_dlp.utils import parse_duration, download_range_func
        start_sec = parse_duration(req.start_time) if req.start_time else 0
        end_sec = parse_duration(req.end_time) if req.end_time else float('inf')
        if start_sec is None or end_sec is None or end_sec <= start_sec:
            raise HTTPException(status_code=400, detail="Rango de tiempo inválido.")
        extra['download_ranges'] = download_range_func(None, [(start_sec, end_sec)])
        extra['force_keyframes_at_cuts'] = True

    def find_output():
        for f in os.listdir(DOWNLOAD_FOLDER):
            if file_uid in f and not f.endswith(('.part', '.ytdl')) and '.temp.' not in f:
                return os.path.join(DOWNLOAD_FOLDER, f)
        return None

    update_progress(req.uid, 5, "Iniciando proceso...")
    try:
        await run_blocking(ytdlp_download, url, extra, lambda: find_output() is not None,
                           on_attempt=lambda n, name: update_progress(req.uid, 5 + n * 5, f"Conectando ({name})..."))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"No se pudo descargar: {str(e)[:300]}")

    path = find_output()
    if not path:
        raise HTTPException(status_code=500, detail="El archivo no apareció después de la descarga.")
    update_progress(req.uid, 100, "¡Archivo listo!")
    _remove_later(background_tasks, path)
    name = os.path.basename(path).replace(f"_{file_uid}", "")
    return FileResponse(path, filename=name)


@app.get("/api/proxy-thumbnail")
async def proxy_thumbnail(url: str):
    """Proxy de miniaturas de Instagram (solo dominios de Instagram/Facebook CDN)."""
    if not is_allowed_thumbnail_url(url):
        raise HTTPException(status_code=400, detail="URL de miniatura no permitida.")
    headers = {'User-Agent': DESKTOP_UA, 'Referer': 'https://www.instagram.com/'}
    try:
        resp = await run_light(requests.get, url, headers=headers, timeout=10)
        resp.raise_for_status()
        ctype = resp.headers.get('Content-Type', 'image/jpeg')
        if not ctype.startswith('image/'):
            return Response(status_code=415)
        return Response(content=resp.content, media_type=ctype)
    except Exception as e:
        logger.debug(f"Proxy de miniatura falló: {e}")
        return Response(status_code=502)


@app.delete("/api/clear-downloads")
async def clear_downloads():
    def clear():
        shutil.rmtree(DOWNLOAD_FOLDER, ignore_errors=True)
        os.makedirs(DOWNLOAD_FOLDER, exist_ok=True)
    try:
        await run_light(clear)
        return {"status": "success", "message": "Descargas locales eliminadas correctamente."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ─── ENDPOINTS: HERRAMIENTAS PERIODÍSTICAS ───────────────────────────────────

@app.post("/api/chat")
async def chat_with_transcript(req: ChatRequest):
    client = require_groq(req.groq_api_key)
    if not req.question.strip():
        raise HTTPException(status_code=400, detail="La pregunta está vacía.")
    try:
        answer = await run_blocking(ai.chat_about_transcript, client, req.transcript, req.question)
        return {"answer": answer}
    except Exception as e:
        raise ai_error(e)


@app.post("/api/analyze")
async def analyze_transcript(req: AnalyzeRequest):
    """Modos: summary (resumen), data (datos duros), angle (ángulos), diarization (hablantes)."""
    client = require_groq(req.groq_api_key)
    if req.mode not in ai.JOURNALIST_PROMPTS:
        raise HTTPException(status_code=400, detail=f"Modo inválido. Opciones: {list(ai.JOURNALIST_PROMPTS)}")
    if len(req.transcript.strip()) < 50:
        raise HTTPException(status_code=400, detail="La transcripción es demasiado corta para analizar.")
    try:
        out = await run_blocking(ai.analyze_transcript, client, req.transcript, req.mode)
        return {"result": out["result"], "mode": req.mode, "model_used": out["model_used"], "parts": out["parts"]}
    except Exception as e:
        raise ai_error(e)


@app.post("/api/quotes")
async def extract_quotes_with_times(req: QuotesRequest):
    """Citas textuales con sus tiempos de entrada/salida en el video."""
    client = require_groq(req.groq_api_key)
    if len(req.transcript.strip()) < 50:
        raise HTTPException(status_code=400, detail="La transcripción es demasiado corta.")
    try:
        quotes, parts = await run_blocking(ai.extract_quotes, client, req.transcript)
    except Exception as e:
        raise ai_error(e)
    if not quotes and parts:
        raise HTTPException(status_code=500, detail="No se pudo interpretar la respuesta de la IA. Reintentá.")

    enriched = []
    for item in quotes:
        quote_text = str(item.get("quote", "")).strip()
        search_kw = str(item.get("search") or quote_text[:60]).strip()
        times = find_segment_times_for_quote(search_kw, req.segments) or find_segment_times_for_quote(quote_text, req.segments)
        enriched.append({"quote": quote_text, "note": str(item.get("note", "")).strip(), "search": search_kw,
                         "start": times.get("start"), "end": times.get("end"), "has_time": bool(times)})
    return {"quotes": enriched, "total": len(enriched), "parts": parts}


@app.post("/api/export-docx")
async def export_docx(req: ExportDocxRequest):
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="Falta la librería python-docx. Ejecutá: pip install python-docx")

    def build() -> io.BytesIO:
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = section.bottom_margin = Inches(1)
            section.left_margin = section.right_margin = Inches(1)
        normal = doc.styles['Normal'].font
        normal.name, normal.size, normal.color.rgb = 'Arial', Pt(11), RGBColor(0x33, 0x41, 0x55)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run(req.title)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(18), RGBColor(0x1e, 0x1b, 0x4b)

        if req.uploader or req.url:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(12)
            for label, value in (("Autor/Canal: ", req.uploader), ("Enlace: ", req.url)):
                if value:
                    p.add_run(label).bold = True
                    p.add_run(f"{value}\n")

        if req.description:
            r = doc.add_paragraph().add_run("Descripción / Copy:")
            r.bold, r.font.size = True, Pt(12)
            p = doc.add_paragraph(req.description)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(18)

        sep = doc.add_paragraph("─" * 40)
        sep.alignment = WD_ALIGN_PARAGRAPH.CENTER

        r = doc.add_paragraph().add_run("Transcripción:")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(14), RGBColor(0x4f, 0x46, 0xe5)

        for para in req.transcript.split("\n\n"):
            para = para.strip()
            if not para:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            if para.startswith("--- ") and para.endswith(" ---"):
                r = p.add_run(para)
                r.bold, r.font.size, r.font.color.rgb = True, Pt(12), RGBColor(0x63, 0x66, 0xf1)
            else:
                # Respetar **negritas** (nombres de hablantes en la diarización)
                for i, piece in enumerate(re.split(r'\*\*(.+?)\*\*', para)):
                    if piece:
                        p.add_run(piece).bold = (i % 2 == 1)
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream

    try:
        stream = await run_light(build)
    except Exception as e:
        logger.exception("Error exportando DOCX")
        raise HTTPException(status_code=500, detail=f"No se pudo generar el archivo DOCX: {e}")
    fname = safe_filename(req.title, 60) or "transcripcion"
    return StreamingResponse(
        stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={'Content-Disposition': f"attachment; filename=\"transcripcion.docx\"; filename*=UTF-8''{quote(fname)}.docx"})


# ─── ENDPOINTS: MANTENIMIENTO ────────────────────────────────────────────────

def _pip_install_requirements() -> tuple:
    req_file = os.path.join(BASE_DIR, 'requirements.txt')
    res = subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', '-r', req_file],
                         capture_output=True, text=True, timeout=900)
    return res.returncode == 0, (res.stderr or res.stdout).strip()[-500:]


def _update_app() -> dict:
    protected = [os.path.join(d, n) for d in (ROOT_DIR, BASE_DIR) for n in (".env", "cookies.txt", "cookies_ig.txt")]
    backups = {}
    for path in protected:
        if os.path.exists(path):
            try:
                with open(path, "rb") as f:
                    backups[path] = f.read()
            except OSError as e:
                logger.warning(f"No se pudo respaldar {path}: {e}")

    repo = os.environ.get('GIT_REPO_DIR', ROOT_DIR)
    error, output = None, ""
    try:
        res = subprocess.run(["git", "pull", "origin", "main"], capture_output=True, text=True, cwd=repo, timeout=300)
        output = res.stdout.strip()
        if res.returncode != 0:
            error = f"git pull salió con código {res.returncode}: {res.stderr.strip()[-400:]}"
    except FileNotFoundError:
        error = "Git no está instalado o no está en el PATH."
    except Exception as e:
        error = str(e)
    finally:
        for path, content in backups.items():
            try:
                with open(path, "wb") as f:
                    f.write(content)
            except OSError as e:
                logger.warning(f"No se pudo restaurar {path}: {e}")

    if error:
        return {"error": f"Error al actualizar: {error}"}
    if "Already up to date" in output or "Ya está actualizado" in output:
        return {"status": "ok", "message": "La aplicación ya está actualizada. No hay cambios nuevos.", "output": output}

    ok, pip_out = _pip_install_requirements()
    msg = "✅ Aplicación actualizada. Reiniciá Clipadsk (cerrá y abrí iniciar.bat) para aplicar los cambios."
    if not ok:
        msg += f" ⚠️ No se pudieron instalar algunas dependencias: {pip_out}"
    return {"status": "ok", "message": msg, "output": output}


@app.post("/api/system/update-app")
async def update_app(request: Request):
    """git pull + reinstalar dependencias, protegiendo .env y cookies."""
    require_admin(request)
    result = await run_light(_update_app)
    if "error" in result:
        return JSONResponse(status_code=500, content=result)
    return result


def _update_engine() -> dict:
    # El backend usa el paquete de Python yt_dlp, así que se actualiza ese (y el .exe si existe)
    res = subprocess.run([sys.executable, '-m', 'pip', 'install', '-U', '-q', 'yt-dlp[default]'],
                         capture_output=True, text=True, timeout=600)
    if res.returncode != 0:
        return {"error": f"Error al actualizar motor: {(res.stderr or res.stdout)[-400:]}"}
    exe = os.path.join(ROOT_DIR, "yt-dlp.exe")
    if os.path.exists(exe):
        try:
            subprocess.run([exe, "-U"], capture_output=True, text=True, timeout=300)
        except Exception as e:
            logger.warning(f"No se pudo actualizar yt-dlp.exe: {e}")
    version = subprocess.run([sys.executable, '-m', 'yt_dlp', '--version'], capture_output=True, text=True).stdout.strip()
    return {"status": "ok",
            "message": f"Motor de descarga actualizado (yt-dlp {version}). Reiniciá Clipadsk para usar la nueva versión.",
            "output": version}


@app.post("/api/system/update-engine")
async def update_engine(request: Request):
    require_admin(request)
    result = await run_light(_update_engine)
    if "error" in result:
        return JSONResponse(status_code=500, content=result)
    return result


@app.post("/api/system/reset")
async def reset_system(request: Request):
    """Borra descargas y la caché de transcripciones."""
    require_admin(request)

    def reset():
        shutil.rmtree(DOWNLOAD_FOLDER, ignore_errors=True)
        os.makedirs(DOWNLOAD_FOLDER, exist_ok=True)
        with db_connect() as conn:
            conn.execute("DELETE FROM transcripts")
    try:
        await run_light(reset)
        return {"status": "ok", "message": "Sistema reseteado (descargas y caché limpias)."}
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": str(e)})


# ─── FRONTEND (debe ir al final para no tapar las rutas /api) ────────────────
if os.path.exists(FRONTEND_DIR):
    FRONTEND_ROOT = Path(FRONTEND_DIR).resolve()
    INDEX_HTML = FRONTEND_ROOT / 'index.html'

    @app.head("/", include_in_schema=False)
    @app.get("/", include_in_schema=False)
    async def serve_index():
        return FileResponse(INDEX_HTML)

    @app.get("/{path:path}", include_in_schema=False)
    async def serve_static_or_index(path: str):
        if path.startswith("api/"):
            raise HTTPException(status_code=404, detail="Endpoint no encontrado")
        try:
            resolved = (FRONTEND_ROOT / path).resolve()
            if FRONTEND_ROOT in resolved.parents and resolved.is_file():
                return FileResponse(resolved)
        except Exception as e:
            logger.debug(f"Error resolviendo ruta estática: {e}")
        return FileResponse(INDEX_HTML)  # rutas de la SPA
else:
    logger.warning(f"No se encontró la carpeta frontend en {FRONTEND_DIR}")


if __name__ == "__main__":
    import uvicorn
    if HOST not in ("127.0.0.1", "localhost") and not ADMIN_TOKEN and not IS_RENDER:
        logger.warning(f"Escuchando en {HOST}: otras PCs de la red pueden usar Clipadsk. Definí ADMIN_TOKEN en .env.")
    logger.info(f"Clipadsk en http://{'127.0.0.1' if HOST == '0.0.0.0' else HOST}:{PORT}")
    uvicorn.run(app, host=HOST, port=PORT, timeout_keep_alive=600)
